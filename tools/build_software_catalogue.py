"""
Build web/software/catalogue.json from the CLG Group Software Catalog workbook.

Usage:  python tools/build_software_catalogue.py "CLG_Group_Software_Catalog_v1_1.xlsx"

Reads the "Master Catalog" sheet (one row per service) and the "Review Queue"
sheet (required actions). Every figure shown in the portal is calculated from
these rows by web/software/software.js, never copied from the workbook's
Summary sheet, so the portal can't drift from the data. To update the portal,
re-run this script on the new workbook and upload the regenerated JSON.
"""
import datetime
import json
import os
import re
import sys

from openpyxl import load_workbook

FIELDS = {
    "Department": "department", "Software / Service": "name", "Vendor": "vendor",
    "Business Purpose": "purpose", "Business Scope": "scope", "Business Owner": "owner",
    "Licence / Cost Model": "licence", "Evidence Source": "evidence",
    "Evidence Confidence": "confidence", "Lifecycle Status": "status", "Review Notes": "notes",
}


def rows(ws):
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h else "" for h in next(it)]
    for r in it:
        if not any(r):
            continue
        yield {header[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(header)}


def slug(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def main(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    items, seen = [], {}
    for r in rows(wb["Master Catalog"]):
        item = {key: r.get(col, "") for col, key in FIELDS.items()}
        if not item["name"]:
            continue
        base = slug(f'{item["department"]}-{item["name"]}')
        seen[base] = seen.get(base, 0) + 1
        item["id"] = base if seen[base] == 1 else f"{base}-{seen[base]}"
        items.append(item)
    actions = {}
    if "Review Queue" in wb.sheetnames:
        for r in rows(wb["Review Queue"]):
            actions[(r.get("Department", ""), r.get("Software / Service", ""))] = r.get("Required Action", "")
    for item in items:
        item["requiredAction"] = actions.get((item["department"], item["name"]), "")
    out = {
        "title": "CLG Group Software Catalog",
        "source": os.path.basename(path),
        "generatedAt": datetime.datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
        "count": len(items),
        "items": items,
    }
    dest = os.path.join(os.path.dirname(__file__), "..", "web", "software", "catalogue.json")
    with open(dest, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print(f"Wrote {len(items)} services to {os.path.normpath(dest)}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "CLG_Group_Software_Catalog_v1_1.xlsx")
